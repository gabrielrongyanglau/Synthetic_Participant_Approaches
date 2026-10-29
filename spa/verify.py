"""Map calculated results to the supplied six-tab workbook for cell-level checks."""
import math
import pandas as pd
from .metrics import DIST,REL

def calculated_cells(b):
    c={}
    def put(sheet,row,col,v):c[(sheet,row,col)]=v
    def vector(s,r,cols,vals):
        for col,v in zip(cols,vals):put(s,r,col,v)
    ss=b['summary'];vs=b['variables'];hs=b['human'];pairs=b['comparisons']
    def summary(method,wave,metric):return ss[(ss.method==method)&(ss.wave==wave)&(ss.metric==metric)].iloc[0]
    for m in range(1,9):
        s=[summary(m,0,k) for k in DIST]
        vector('Goal 1',6+m,range(2,9),[s[0]['n']]+[x['mean'] for x in s]+[s[3].lower,s[3].upper])
        for k,start in [(2,7),(3,20)]:
            s=[summary(m,k,f) for f in REL]
            vector('Goal 2',start+m,range(2,9),[s[0]['n']]+[x['mean'] for x in s]+[s[0].lower,s[0].upper])
    from .study import load_study
    ps=load_study()['psychological_predictors']
    for j,code in enumerate(ps):
        h=hs[hs.code==code].iloc[0]
        vector('By psychological variable',11+j*13,[2,4,6,9,12],[h['n'],h['mean'],h['sd'],h.midus2_r,h.midus3_r])
        for m in range(1,9):
            v=vs[(vs.method==m)&(vs.code==code)].iloc[0]
            cols=['distribution_runs','mean','sd','smd','sd_ratio','sd_error','wasserstein','midus2_runs','midus2_mean_r','midus2_mae','midus3_runs','midus3_mean_r','midus3_mae']
            vector('By psychological variable',13+j*13+m,range(2,15),[v[x] for x in cols])
    families=[(0,m) for m in DIST]+[(k,m) for k in (2,3) for m in REL]
    for i,(wave,metric) in enumerate(families):
        for m in range(1,9):
            s=summary(m,wave,metric)
            vector('95% CIs',9+i*8+m,range(5,9),[s['n'],s['mean'],s.lower,s.upper])
        p=pairs[(pairs.wave==wave)&(pairs.metric==metric)].sort_values(['first','second'])
        for j,(_,x) in enumerate(p.iterrows()):
            pval='<1e-300' if x.p_adjusted<1e-300 else x.p_adjusted
            vector('Method comparisons',11+i*28+j,range(5,10),[x.difference,x.lower,x.upper,pval,'Yes' if x.significant else 'No'])
    return c

def verify(b,workbook,output):
    from openpyxl import load_workbook
    from pathlib import Path
    wb=load_workbook(workbook,data_only=True);cells=calculated_cells(b);rows=[]
    for s in wb:
        for row in s:
            for x in row:
                key=(s.title,x.row,x.column)
                if isinstance(x.value,(int,float)) and key not in cells:raise ValueError('Unmapped numeric reference cell '+str(key))
    for (sheet,r,col),v in cells.items():
        target=wb[sheet].cell(r,col).value
        if isinstance(target,(int,float)) and isinstance(v,(int,float)):
            d=abs(float(target)-float(v));is_p=sheet=='Method comparisons' and col==8
            count=(sheet in ('Goal 1','Goal 2') and col==2) or (sheet=='95% CIs' and col==5) or (sheet=='By psychological variable' and ((r-11)%13==0 and col==2 or (r-14)%13 in range(8) and col in (2,9,12)))
            tol=0 if count else max(1e-310,abs(target)*1e-7) if is_p else 1e-10
            ok=math.isfinite(float(v)) and d<=tol
        else:d=0 if target==v else math.inf;ok=target==v;tol=0
        rows.append(dict(sheet=sheet,row=r,column=col,reference=target,calculated=v,absolute_difference=d,tolerance=tol,passed=ok))
    frame=pd.DataFrame(rows);root=Path(output);root.mkdir(parents=True,exist_ok=True)
    frame.to_csv(root/'validation_cells.csv',index=False)
    frame.groupby('sheet').agg(checked=('passed','size'),passed=('passed','sum'),maximum_difference=('absolute_difference','max')).to_csv(root/'validation_summary.csv')
    failures=frame[~frame.passed]
    if len(failures):raise ValueError(f'{len(failures)} reference discrepancies. See {root}/validation_cells.csv')
    return {'checked':len(frame),'passed':int(frame.passed.sum()),'maximum_absolute_difference':float(frame.absolute_difference.max())}
